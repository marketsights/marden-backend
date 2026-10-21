"""Document text for the agents: one text per document with page boundaries, plus quote finding.

    doc = read_text(data, content_type, url)        # DocText(kind, text, pages, scanned, ...)
    hit = find_quote(doc, quote)                    # QuoteHit(start, end, page, mode) or None

Supported: PDF (PyMuPDF; pages from the file), HTML (SEC filings: page breaks from page-break styles;
hidden iXBRL headers skipped), DOCX (FERC issuances; explicit page breaks), PPTX (one page per slide),
XLSX (one page per sheet), plain text / CSV / JSON. Pages are separated by form feeds in `text`.

The quote check is layer 1 of spec 22 ("no quote, no fact"): a quote must appear verbatim after
whitespace normalization. Normalization also folds what is rendering, not wording: Unicode compatibility
forms (ligatures), typographic quotes and dashes, non-breaking and zero-width spaces, soft hyphens, and
(mode 'dehyphenated') a word broken across a PDF line ("con- tract").
"""

from __future__ import annotations

import io
import re
import unicodedata
import zipfile
from bisect import bisect_right
from dataclasses import dataclass, field

from lxml import etree
from lxml import html as lhtml

FORM_FEED = "\f"
_WS = re.compile(r"\s+")
_FOLD = str.maketrans({
    **dict.fromkeys(["\u2018", "\u2019", "\u201a", "\u201b", "\u2032"], "'"),
    **dict.fromkeys(["\u201c", "\u201d", "\u201e", "\u2033"], '"'),
    **dict.fromkeys(["\u2010", "\u2011", "\u2012", "\u2013", "\u2014", "\u2015", "\u2212"], "-"),
    **dict.fromkeys(["\u00a0", "\u2007", "\u202f", "\u2009"], " "),
    **dict.fromkeys(["\u00ad", "\u200b", "\u200c", "\u200d", "\ufeff"], None),
})
_DEHYPHEN = re.compile(r"(?<=[a-z])- (?=[a-z])")
BLOCK_TAGS = frozenset({
    "p", "div", "br", "tr", "li", "ul", "ol", "table", "h1", "h2", "h3", "h4", "h5", "h6", "section",
    "article", "header", "footer", "blockquote", "pre", "hr", "dt", "dd", "center", "title",
})
CELL_TAGS = frozenset({"td", "th"})
SKIP_TAGS = frozenset({"script", "style", "head", "noscript", "template"})
_PAGE_BREAK = re.compile(r"page-break-(before|after)\s*:\s*always|break-(before|after)\s*:\s*page", re.I)
_HIDDEN = re.compile(r"display\s*:\s*none", re.I)


@dataclass
class DocText:
    kind: str
    text: str
    pages: list[tuple[int, int]]  # [start, end) char span of each page in `text`
    scanned: bool = False
    title: str | None = None
    _norm: tuple[str, list[int]] | None = field(default=None, repr=False)
    _dehy: tuple[str, list[int]] | None = field(default=None, repr=False)

    @property
    def page_count(self) -> int:
        return len(self.pages)

    def page_of(self, offset: int) -> int:
        """1-based page number holding a char offset."""
        starts = [p[0] for p in self.pages]
        return max(1, bisect_right(starts, offset))

    def page_text(self, first: int, last: int) -> str:
        """Text of pages first..last (1-based, inclusive)."""
        first, last = max(1, first), min(self.page_count, last)
        if first > last:
            return ""
        return self.text[self.pages[first - 1][0]: self.pages[last - 1][1]]


@dataclass(frozen=True)
class QuoteHit:
    start: int
    end: int
    page: int
    mode: str  # 'exact' (normalized whitespace) or 'dehyphenated'


# ------------------------------------------------------------------ normalization and quotes
def normalize(text: str) -> str:
    return _WS.sub(" ", unicodedata.normalize("NFKC", text).translate(_FOLD)).strip()


def _normalized_with_map(text: str) -> tuple[str, list[int]]:
    """Normalized text and, per normalized char, the offset of the original char it came from."""
    out: list[str] = []
    index: list[int] = []
    pending_space = False
    for i, ch in enumerate(text):
        piece = unicodedata.normalize("NFKC", ch).translate(_FOLD)
        for c in piece:
            if c.isspace():
                pending_space = bool(out)
                continue
            if pending_space:
                out.append(" ")
                index.append(i)
                pending_space = False
            out.append(c)
            index.append(i)
    return "".join(out), index


def _dehyphenated(norm: str, index: list[int]) -> tuple[str, list[int]]:
    out: list[str] = []
    idx: list[int] = []
    last = 0
    for m in _DEHYPHEN.finditer(norm):
        out.append(norm[last:m.start()])
        idx.extend(index[last:m.start()])
        last = m.end()
    out.append(norm[last:])
    idx.extend(index[last:])
    return "".join(out), idx


def find_quote(doc: DocText, quote: str) -> QuoteHit | None:
    q = normalize(quote or "")
    if len(q) < 3:
        return None
    if doc._norm is None:
        doc._norm = _normalized_with_map(doc.text)
    norm, index = doc._norm
    pos = norm.find(q)
    mode = "exact"
    if pos < 0:
        if doc._dehy is None:
            doc._dehy = _dehyphenated(norm, index)
        norm, index = doc._dehy
        pos = norm.find(_DEHYPHEN.sub("", q))
        mode = "dehyphenated"
        q = _DEHYPHEN.sub("", q)
    if pos < 0:
        return None
    start, end = index[pos], index[pos + len(q) - 1] + 1
    return QuoteHit(start, end, doc.page_of(start), mode)


def context(doc: DocText, hit: QuoteHit, chars: int = 1500) -> str:
    """The passage around a quote (for the verifier), cut at whitespace."""
    lo, hi = max(0, hit.start - chars), min(len(doc.text), hit.end + chars)
    while lo > 0 and not doc.text[lo - 1].isspace() and hit.start - lo < chars + 80:
        lo -= 1
    while hi < len(doc.text) and not doc.text[hi].isspace() and hi - hit.end < chars + 80:
        hi += 1
    return doc.text[lo:hi].replace(FORM_FEED, "\n")


# ------------------------------------------------------------------ readers
def _join_pages(pages: list[str]) -> tuple[str, list[tuple[int, int]]]:
    spans: list[tuple[int, int]] = []
    parts: list[str] = []
    pos = 0
    for i, page in enumerate(pages):
        if i:
            parts.append(FORM_FEED)
            pos += 1
        spans.append((pos, pos + len(page)))
        parts.append(page)
        pos += len(page)
    return "".join(parts), spans or [(0, 0)]


def _tidy(text: str) -> str:
    lines = [re.sub(r"[ \t\u00a0]+", " ", ln).strip() for ln in text.split("\n")]
    out = re.sub(r"\n{3,}", "\n\n", "\n".join(lines))
    return out.strip()


def read_pdf(data: bytes) -> DocText:
    import pymupdf

    with pymupdf.open(stream=data, filetype="pdf") as pdf:
        pages = [_tidy(page.get_text("text")) for page in pdf]
        images = sum(len(page.get_images()) for page in pdf)
        title = (pdf.metadata or {}).get("title") or None
    text, spans = _join_pages(pages)
    chars = sum(len(p) for p in pages)
    # Scanned: page images and (almost) no text layer. Short text pages without images are just short.
    scanned = bool(pages) and images > 0 and chars < 40 * len(pages)
    return DocText("pdf", text, spans, scanned=scanned, title=title)


def _break(style: str, where: str) -> bool:
    m = _PAGE_BREAK.search(style)
    return bool(m and (m.group(1) or m.group(2) or "").lower() == where)


def _html_pages(root: etree._Element) -> list[list[str]]:
    """Text pieces per page; iterative (filings nest deeply), skipping hidden and non-content elements."""
    pages: list[list[str]] = [[]]
    skip_depth = 0
    for event, el in etree.iterwalk(root, events=("start", "end")):
        tag = el.tag.lower() if isinstance(el.tag, str) else ""
        hidden = not tag or tag in SKIP_TAGS or ":header" in tag or bool(_HIDDEN.search(el.get("style") or ""))
        if event == "start":
            if skip_depth or hidden:
                skip_depth += 1
                continue
            style = el.get("style") or ""
            if _break(style, "before") and any(x.strip() for x in pages[-1]):
                pages.append([])
            if tag in BLOCK_TAGS:
                pages[-1].append("\n")
            if el.text:
                pages[-1].append(el.text)
            continue
        if skip_depth:
            skip_depth -= 1
            if skip_depth == 0 and el.tail:
                pages[-1].append(el.tail)
            continue
        if tag in CELL_TAGS:
            pages[-1].append("\t")
        if tag in BLOCK_TAGS:
            pages[-1].append("\n")
        if _break(el.get("style") or "", "after"):
            pages.append([])
        if el.tail:
            pages[-1].append(el.tail)
    return pages


def read_html(data: bytes) -> DocText:
    root = lhtml.fromstring(data)
    title_el = root.find(".//title")
    body = root.find("body")
    pages = _html_pages(body if body is not None else root)
    texts = [_tidy("".join(p)) for p in pages]
    texts = [t for t in texts if t] or [""]
    text, spans = _join_pages(texts)
    title = title_el.text_content().strip() if title_el is not None and title_el.text_content() else None
    return DocText("html", text, spans, title=title)


_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def read_docx(data: bytes) -> DocText:
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        root = etree.fromstring(z.read("word/document.xml"))
    pages: list[list[str]] = [[]]
    for el in root.iter():
        if el.tag == f"{_W}p":
            pages[-1].append("\n")
        elif el.tag == f"{_W}t" and el.text:
            pages[-1].append(el.text)
        elif el.tag == f"{_W}tab" or el.tag == f"{_W}tc":
            pages[-1].append("\t")
        elif el.tag == f"{_W}br" and el.get(f"{_W}type") == "page":
            pages.append([])
        elif el.tag == f"{_W}br":
            pages[-1].append("\n")
    texts = [t for t in (_tidy("".join(p)) for p in pages) if t] or [""]
    text, spans = _join_pages(texts)
    return DocText("docx", text, spans)


def read_pptx(data: bytes) -> DocText:
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        names = [n for n in z.namelist() if re.fullmatch(r"ppt/slides/slide\d+\.xml", n)]
        names.sort(key=lambda n: int(re.search(r"(\d+)\.xml$", n).group(1)))  # type: ignore[union-attr]
        slides = []
        for name in names:
            root = etree.fromstring(z.read(name))
            paras = ["".join((t.text or "") if t.tag == f"{_A}t" else "\n" for t in p.iter(f"{_A}t", f"{_A}br"))
                     for p in root.iter(f"{_A}p")]
            slides.append(_tidy("\n".join(paras)))
    text, spans = _join_pages(slides or [""])
    return DocText("pptx", text, spans)


def read_xlsx(data: bytes, max_rows: int = 20_000) -> DocText:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sheets = []
    try:
        for ws in wb.worksheets:
            rows = [f"# {ws.title}"]
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i >= max_rows:
                    rows.append(f"... ({ws.title}: rows after {max_rows} not read)")
                    break
                cells = ["" if v is None else str(v) for v in row]
                if any(cells):
                    rows.append("\t".join(cells).rstrip())
            sheets.append("\n".join(rows))
    finally:
        wb.close()
    text, spans = _join_pages(sheets or [""])
    return DocText("xlsx", text, spans)


def read_plain(data: bytes) -> DocText:
    raw = data.decode("utf-8", errors="replace")
    pages = [_tidy(p) for p in raw.split(FORM_FEED)]
    text, spans = _join_pages(pages)
    return DocText("txt", text, spans)


def sniff(data: bytes, content_type: str | None = None, url: str | None = None) -> str:
    """'pdf' | 'docx' | 'pptx' | 'xlsx' | 'html' | 'txt' | 'unknown', from the bytes first."""
    head = data[:2048].lstrip()
    if data[:5] == b"%PDF-":
        return "pdf"
    if data[:2] == b"PK":
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                names = set(z.namelist())
        except zipfile.BadZipFile:
            return "unknown"
        if "word/document.xml" in names:
            return "docx"
        if any(n.startswith("ppt/slides/") for n in names):
            return "pptx"
        if "xl/workbook.xml" in names:
            return "xlsx"
        return "unknown"
    low = head[:1024].lower()
    if low.startswith((b"<!doctype html", b"<html", b"<?xml")) and b"<html" in data[:4096].lower():
        return "html"
    if (content_type or "").lower().startswith(("text/html", "application/xhtml")):
        return "html"
    try:
        data[:65536].decode("utf-8")
    except UnicodeDecodeError:
        return "unknown"
    return "txt"


READERS = {"pdf": read_pdf, "html": read_html, "docx": read_docx, "pptx": read_pptx, "xlsx": read_xlsx,
           "txt": read_plain}


class UnreadableDocument(ValueError):
    pass


def read_text(data: bytes, content_type: str | None = None, url: str | None = None) -> DocText:
    kind = sniff(data, content_type, url)
    reader = READERS.get(kind)
    if reader is None:
        raise UnreadableDocument(f"unsupported document type ({content_type or 'unknown'})")
    try:
        return reader(data)
    except (zipfile.BadZipFile, KeyError, etree.XMLSyntaxError, ValueError, RuntimeError) as exc:
        raise UnreadableDocument(f"{kind} could not be read: {type(exc).__name__}") from None


def pdf_pages(data: bytes, first: int, last: int) -> bytes:
    """A new PDF holding pages first..last (1-based, inclusive) of `data` (a section's own document)."""
    import pymupdf

    with pymupdf.open(stream=data, filetype="pdf") as src, pymupdf.open() as out:
        out.insert_pdf(src, from_page=first - 1, to_page=last - 1)
        return out.tobytes(garbage=3, deflate=True)
