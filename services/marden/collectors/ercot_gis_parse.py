"""Parse ERCOT's monthly GIS Report (EMIL PG7-200-ER, "GIS_Report_<Month><Year>.xlsx") into queue entries.

Layout (the workbook as ERCOT posts it; the same sheets and header names gridstatus reads):
- "Project Details - Large Gen" and "Project Details - Small Gen": one row per generation interconnection
  request (INR), header row found by its "INR" cell (a few sub-header rows follow it, then data). Named
  columns used: INR, Project Name, Interconnecting Entity, POI Location, County, CDR Reporting Zone,
  Projected COD, Fuel, Technology, Capacity (MW), GIM Study Phase, Screening Study Started, IA Signed,
  Approved for Energization, Approved for Synchronization.
- "Inactive Projects" (Inactive Date) and "Cancellation Update" (Cancel Date): INR, Project Name, County,
  Fuel, "MW **", Size Category.
Rows whose INR cell is not an INR number (sub-headers, notes) are skipped. A project listed on several
sheets keeps the last one in the order above (a cancellation is the newest record of a request). INR
numbers keep ERCOT's form: the year/INR/number upper case, a split-project suffix lower case ('15INR0064b',
as ERCOT and LBNL write it).

The MIS listing (https://www.ercot.com/misapp/GetReports.do?reportTypeId=15933, `mis_listing`): an HTML
table, one row per posted file, its title "RPT.00015933.0000000000000000.<YYYYMMDD>.<HHMMSSmmm>.<file>"
(the posting time, ERCOT local) and one link to /misdownload/servlets/mirDownload?...doclookupId=<id>.
The same report type also carries "Co-located_Battery_Identification_Report_<Month>_<Year>.xlsx" files.

Status: the GIS report has no status column, so `status` is the latest milestone ERCOT reports for the
request, in ERCOT's column words: 'Approved for Synchronization', 'Approved for Energization', 'IA
Signed', else the GIM Study Phase; 'Inactive' and 'Cancelled' on those sheets. `study_phase` is the GIM
Study Phase plus the milestone dates. ERCOT removes a project
from the GIS report when it reaches commercial operation, so nothing here is 'in_service'.
State: ERCOT's grid lies within Texas and the report has no state column, so every entry is TX.
queue_date: "Screening Study Started" (the report does not publish the request date).
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass
from datetime import date, datetime
from html.parser import HTMLParser
from typing import Any
from urllib.parse import parse_qs, urlsplit

import openpyxl

from marden.collectors.iso_queue import QueueEntry, month_or_date, number, text, us_date
from marden.collectors.rates_common import ParseError

LARGE = "Project Details - Large Gen"
SMALL = "Project Details - Small Gen"
INACTIVE = "Inactive Projects"
CANCELLED = "Cancellation Update"
SHEETS = (LARGE, SMALL, INACTIVE, CANCELLED)
REQUIRED = (LARGE,)
# Fuel and technology codes (the report's acronym list).
FUEL = {
    "BIO": "Biomass", "COA": "Coal", "GAS": "Gas", "GEO": "Geothermal", "HYD": "Hydrogen", "NUC": "Nuclear",
    "OIL": "Fuel Oil", "OTH": "Other", "PET": "Petcoke", "SOL": "Solar", "WAT": "Water", "WIN": "Wind",
}
TECHNOLOGY = {
    "BA": "Battery Energy Storage", "CC": "Combined-Cycle", "CE": "Compressed Air Energy Storage",
    "CP": "Concentrated Solar Power", "EN": "Energy Storage", "FC": "Fuel Cell",
    "GT": "Combustion (gas) Turbine, but not part of a Combined-Cycle", "HY": "Hydroelectric Turbine",
    "IC": "Internal Combustion Engine, eg. Reciprocating", "OT": "Other", "PV": "Photovoltaic Solar",
    "ST": "Steam Turbine other than Combined-Cycle", "WT": "Wind Turbine",
}
COLUMNS = {  # field -> published header names (whitespace collapsed, lower case)
    "inr": ("inr",),
    "name": ("project name",),
    "developer": ("interconnecting entity",),
    "poi": ("poi location",),
    "county": ("county",),
    "zone": ("cdr reporting zone",),
    "cod": ("projected cod",),
    "fuel": ("fuel",),
    "technology": ("technology",),
    "mw": ("capacity (mw)", "mw **", "mw"),
    "phase": ("gim study phase",),
    "screening": ("screening study started",),
    "ia": ("ia signed",),
    "energization": ("approved for energization",),
    "sync": ("approved for synchronization",),
    "size": ("size category",),
    "inactive": ("inactive date",),
    "cancel": ("cancel date",),
}
_INR = re.compile(r"^(\d{2}INR\d{3,5})([A-Z]?)$", re.I)
MAX_HEADER_ROW = 60


def _norm(value: Any) -> str:
    return " ".join(str(value).split()).lower() if value is not None else ""


def _inr(value: Any) -> str | None:
    m = _INR.match((text(value) or "").replace(" ", ""))
    return m[1].upper() + m[2].lower() if m else None


def _header(name: str, rows: list[tuple[Any, ...]]) -> tuple[int, dict[str, int]]:
    for i, row in enumerate(rows[:MAX_HEADER_ROW]):
        names = [_norm(c) for c in row]
        if "inr" in names:
            cols = {}
            for fld, options in COLUMNS.items():
                idx = next((names.index(o) for o in options if o in names), None)
                if idx is not None:
                    cols[fld] = idx
            return i + 1, cols
    raise ParseError(f"GIS sheet {name!r} has no 'INR' header row in its first {MAX_HEADER_ROW} rows")


def _code(mapping: dict[str, str], value: Any) -> str | None:
    s = text(value)
    return mapping.get(s.upper(), s) if s else None


def parse_sheet(name: str, rows: list[tuple[Any, ...]]) -> list[QueueEntry]:
    start, cols = _header(name, rows)
    required = ("inr", "name", "county") + (("cod", "fuel") if name in (LARGE, SMALL) else ())
    missing = [f for f in required if f not in cols]
    if missing:
        raise ParseError(f"GIS sheet {name!r} lacks columns {missing}")

    def cell(row: tuple[Any, ...], fld: str) -> Any:
        i = cols.get(fld)
        return row[i] if i is not None and i < len(row) else None

    out = []
    for row in rows[start:]:
        inr = _inr(cell(row, "inr"))
        if inr is None:
            continue
        fuel, tech = _code(FUEL, cell(row, "fuel")), _code(TECHNOLOGY, cell(row, "technology"))
        size = {LARGE: "Large", SMALL: "Small"}.get(name) or text(cell(row, "size"))
        phase = text(cell(row, "phase"))
        values: dict[str, Any] = {
            "name": text(cell(row, "name")),
            "developer": text(cell(row, "developer")),
            "county": text(cell(row, "county")),
            "state": "TX",
            "poi": text(cell(row, "poi")),
            "fuel": fuel,
            "project_type": " ".join(p for p in (f"{size} Gen" if size else "Generator", f"({tech})" if tech else "")
                                     if p),
            "mw_requested": number(cell(row, "mw")),
            "queue_date": us_date(cell(row, "screening")),
            "proposed_cod": month_or_date(cell(row, "cod")),
            "zone": text(cell(row, "zone")),
        }
        if name == CANCELLED:
            values.update(status="Cancelled", status_class="withdrawn", withdrawn_date=us_date(cell(row, "cancel")))
        elif name == INACTIVE:
            since = us_date(cell(row, "inactive"))
            values.update(status="Inactive", status_class="suspended",
                          study_phase=f"Inactive since {since.isoformat()}" if since else None)
        else:
            sync, energized, ia = (us_date(cell(row, f)) for f in ("sync", "energization", "ia"))
            if sync:
                status, cls = "Approved for Synchronization", "under_construction"
            elif energized:
                status, cls = "Approved for Energization", "under_construction"
            elif ia:
                status, cls = "IA Signed", "active"
            else:
                status, cls = phase or "In study", "active"
            dates = [f"{label} {d.isoformat()}" for label, d in (
                ("IA signed", ia), ("approved for energization", energized), ("approved for synchronization", sync))
                if d]
            values.update(status=status, status_class=cls, study_phase="; ".join(([phase] if phase else []) + dates)
                          or None)
        out.append(QueueEntry(inr, values))
    return out


def parse_workbook(data: bytes) -> list[QueueEntry]:
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises zipfile/KeyError/InvalidFileException for non-workbooks
        raise ParseError(f"ERCOT GIS file is not an xlsx workbook: {type(exc).__name__}: {exc}") from None
    by_name = {" ".join(n.split()).lower(): n for n in wb.sheetnames}
    missing = [s for s in REQUIRED if s.lower() not in by_name]
    if missing:
        raise ParseError(f"ERCOT GIS workbook lacks sheets {missing}; has {wb.sheetnames}")
    entries: dict[str, QueueEntry] = {}
    for sheet in SHEETS:
        if sheet.lower() in by_name:
            for e in parse_sheet(sheet, list(wb[by_name[sheet.lower()]].iter_rows(values_only=True))):
                entries[e.queue_id] = e
    return list(entries.values())


_MONTHS = ("january", "february", "march", "april", "may", "june", "july", "august", "september", "october",
           "november", "december")


def report_month(friendly_name: str) -> date | None:
    """'GIS_Report_September2026' -> 2026-09-01 (the month the report covers); ERCOT also writes
    'GIS_Report_Jun2026' and 'GIS_Report_May_2020'."""
    m = re.search(r"GIS_Report_([A-Za-z]+)_?(\d{4})", friendly_name, re.I)
    word = m[1].lower() if m else ""
    month = next((i for i, name in enumerate(_MONTHS, 1) if len(word) >= 3 and name.startswith(word)), None)
    return date(int(m[2]), month, 1) if m and month else None


# ---------------------------------------------------------------- MIS listing
_POSTING = re.compile(r"^RPT\.(\d+)\.\d+\.(\d{8})\.(\d{6})\d*\.(.+)$")


@dataclass(frozen=True)
class Posting:
    """One file on the MIS listing."""

    doc_id: str  # doclookupId
    file: str  # 'GIS_Report_September2026.xlsx'
    posted_at: datetime  # ERCOT's posting time, from the listed title (local time, naive)
    report_type: str  # '15933'

    @property
    def name(self) -> str:
        return self.file.rsplit(".", 1)[0]

    @property
    def posted(self) -> date:
        return self.posted_at.date()

    @property
    def month(self) -> date | None:
        return report_month(self.file)

    @property
    def is_gis(self) -> bool:
        return self.file.lower().startswith("gis_report") and self.file.lower().endswith(".xlsx")


class _Rows(HTMLParser):
    """Table rows as (texts, hrefs); a nested table's rows are their own rows."""

    def __init__(self) -> None:
        super().__init__()
        self.stack: list[tuple[list[str], list[str]]] = []
        self.rows: list[tuple[list[str], list[str]]] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag == "tr":
            self.stack.append(([], []))
        elif tag == "a" and self.stack:
            self.stack[-1][1].extend(v for k, v in attrs if k == "href" and v)

    def handle_data(self, data: str) -> None:
        if self.stack and data.strip():
            self.stack[-1][0].append(data.strip())

    def handle_endtag(self, tag: str) -> None:
        if tag == "tr" and self.stack:
            self.rows.append(self.stack.pop())


def mis_listing(html: bytes) -> list[Posting]:
    """Every posting on a MIS GetReports.do page, as listed (newest first)."""
    parser = _Rows()
    parser.feed(html.decode("utf-8", errors="replace"))
    out = []
    for texts, hrefs in parser.rows:
        title = next((m for t in texts if (m := _POSTING.match(t))), None)
        doc_id = next((q["doclookupId"][0] for h in hrefs if "doclookupId" in (q := parse_qs(urlsplit(h).query))),
                      None)
        if title is None or doc_id is None or not doc_id.isdigit():
            continue
        try:
            posted_at = datetime.strptime(title[2] + title[3], "%Y%m%d%H%M%S")
        except ValueError:
            raise ParseError(f"MIS posting {title[0]!r} has no valid posting time") from None
        out.append(Posting(doc_id, title[4], posted_at, str(int(title[1]))))
    return out
