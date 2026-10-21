"""NYISO interconnection queue -> queue_project facts (tag F, source NYISO). See iso_queue for the model.

NYISO posts one workbook a month, linked as "Interconnection Queue" on
https://www.nyiso.com/interconnections (e.g. NYISO-Interconnection-Queue-09302026.xlsx: the queue as of
Sep 30, 2026). The page is stored raw, the current workbook link is read from it, and the workbook is
stored raw before parsing. Sheets used: Interconnection Queue, Cluster Projects, Load Projects (large
loads: data centers, crypto mining, fabs, hydrogen; fuel LOAD), In Service, Withdrawn and Cluster
Projects-Withdrawn. Affected System Studies (other ISOs' projects) and the summary sheets are skipped.
Status codes and fuel codes are translated with the keys printed in the workbook's own notes.

Cadence: monthly publication; checked weekly (Wednesday) so a new month is picked up within days.
`python -m marden collect nyiso_queue`.
"""

from __future__ import annotations

import argparse
import io
import re
from datetime import date
from typing import Any
from urllib.parse import urljoin

import openpyxl

from marden import cli
from marden.collectors.iso_queue import LOAD, QueueEntry, QueueStore, month_or_date, number, text, us_date
from marden.collectors.rates_common import ET, ParseError, RunStats, collect, print_stats
from marden.core import db, http
from marden.core.documents import DocumentRow
from marden.core.http import Fetcher, fetch_document
from marden.core.queue import app
from marden.core.sources import declare_source

SOURCE_ID = "nyiso_queue"
PRODUCER = "collector:nyiso_queue"
HOST = "www.nyiso.com"
INDEX_URL = f"https://{HOST}/interconnections"
CRON_UTC = "37 11 * * 3"  # Wednesdays 11:37 UTC
TERMS_URL = "https://www.nyiso.com/legal-notice"
TERMS_NOTE = (
    f"Checked 2026-10-21. NYISO Legal Notice ({TERMS_URL}): \"Access to this Web site does not confer any license "
    "or ownership interest in either the form or content of the Web site ... Downloading, republishing, "
    "retransmitting, reproducing, or other use of any image or video on this website as a stand-alone file is "
    "strictly prohibited\" (images and video only; no clause on data files or automated access). robots.txt "
    "(https://www.nyiso.com/robots.txt) disallows */search, /o/, /group/, /content-not-found/, /public/ and "
    "/viewassets; /interconnections and /documents/ are allowed. The workbook is the public 'Interconnection "
    "Queue' download (no MyNYISO account). One page and one workbook a week."
)
SOURCE = declare_source(SOURCE_ID, "NYISO", "portal", f"https://{HOST}", 7 * 24 * 60, TERMS_NOTE)
http.set_rate_limit(HOST, 0.5)

_LINK = re.compile(r'href="(/documents/[^"]*?/NYISO-Interconnection-Queue-(\d{8})\.xlsx[^"]*)"', re.I)
# 'Project Status # Key' notes in the Interconnection Queue / Load Projects sheets.
STATUS = {
    "0": "Withdrawn", "1": "Scoping Meeting Pending", "2": "FES Pending", "3": "FES in Progress",
    "3A": "FES Approved/Performed", "4": "SRIS/SIS Pending", "5": "SRIS/SIS in Progress",
    "5P": "SRIS Commenced, Stopped and Pending Adoption of IP", "6": "SRIS/SIS Approved", "7": "FS Pending",
    "8": "Rejected Cost Allocation/Next FS Pending", "9": "FS in Progress",
    "10": "Accepted Cost Allocation/IA in Progress", "11": "IA Completed", "12": "Under Construction",
    "13": "In Service for Test", "14": "In Service Commercial", "15": "Partial In-Service",
    "P": "Pending Adoption of IP Compliance with Order 2023",
}
# 'Project Status # Key' note in the Cluster Projects sheet (codes are published without the C).
CLUSTER_STATUS = {
    "1": "IR Validated/ Scoping Meeting Pending", "2": "Customer Engagement Window",
    "3": "Phase 1 Entry Decision Period", "4": "Phase 1 Study", "5": "Phase 2 Entry Decision Period",
    "6": "Phase 2 Study", "7": "Final Decision Period", "10": "Accepted Cost Allocation/ IA in Progress",
    "11": "IA Completed", "12": "Under Construction", "13": "In Service for Test", "14": "In Service Commercial",
    "15": "Partially In-Service/ Partially under construction", "0": "Withdrawn",
}
# 'Type / Fuel Key' note (codes the note lists without a name are kept as published).
FUEL = {
    "ST": "Steam Turbine", "CT": "Combustion Turbine", "CS": "Steam Turbine & Combustion Turbine", "H": "Hydro",
    "PS": "Pumped Storage", "W": "Wind", "OSW": "Off-Shore Wind", "NU": "Nuclear", "NG": "Natural Gas",
    "M": "Methane", "SW": "Solid Waste", "S": "Solar", "WO": "Wood", "F": "Flywheel", "ES": "Energy Storage",
    "O": "Oil", "C": "Coal", "D": "Dual Fuel", "L": LOAD, "FC": "Fuel Cell", "CW": "CSR - ES + Wind",
    "CR": "CSR - ES + Solar",
}
# 'End-Use Key' note in the Load Projects sheet.
END_USE = {
    "DAT": "Data Center", "DAT-AI": "AI Data center", "DAT-CM": "Cryptocurrency mining data center",
    "M": "Manufacturing", "M-CH": "Microchip fabrication", "M-CG": "Consumable Goods", "M-IN": "Industrial Goods",
    "RD": "Research and Development", "HP": "Hydrogen Production", "O": "Other",
}
# sheet -> (status key, status when the cell is blank)
SHEETS = {
    "Interconnection Queue": (STATUS, None),
    "Cluster Projects": (CLUSTER_STATUS, None),
    "Load Projects": (STATUS, None),
    "In Service": (STATUS, "In Service"),
    "Withdrawn": (STATUS, "Withdrawn"),
    "Cluster Projects-Withdrawn": (CLUSTER_STATUS, "Withdrawn"),
}
COLUMNS = {  # field -> header names across the sheets (lower case, spaces collapsed)
    "queue_id": ("queue pos.", "queue number"),
    "developer": ("developer/interconnection customer", "interconnection customer name", "developer name",
                  "owner/developer"),
    "name": ("project name", "project: project name"),
    "queue_date": ("date of ir", "ir submission date"),
    "mw": ("sp (mw)", "peak mw load"),
    "fuel": ("type/ fuel", "type/fuel"),
    "county": ("county", "location county"),
    "state": ("state",),
    "zone": ("nyiso zone", "z", "zone"),
    "poi": ("points of interconnection", "interconnection point"),
    "utility": ("utility", "cto/utility"),
    "status": ("project status #", "s"),
    "updated": ("last updated date", "last update"),
    "in_service": ("proposed in-service/initial backfeed date", "proposed initial backfeed date",
                   "proposed in-service"),
    "cod": ("proposed cod",),
    "end_use": ("end-use",),
    "record_type": ("record type", "record type name", "project type"),
    "studies": ("availability of studies",),
}
_QID = re.compile(r"^(?=.*\d)[A-Z0-9][A-Z0-9-]{0,19}$")  # 0276, 0290A, C24-003, C26-0001-1


def _norm(value: Any) -> str:
    return " ".join(str(value).split()).lower() if value is not None else ""


def find_workbook(page: bytes) -> tuple[str, date]:
    """The workbook URL and its as-of date (from the file name) on the interconnections page."""
    m = _LINK.search(page.decode("utf-8", errors="replace"))
    if not m:
        raise ParseError(f"no NYISO-Interconnection-Queue-MMDDYYYY.xlsx link on {INDEX_URL}")
    stamp = m[2]
    try:
        as_of = date(int(stamp[4:]), int(stamp[:2]), int(stamp[2:4]))
    except ValueError:
        raise ParseError(f"bad date in workbook name: {stamp}") from None
    return urljoin(INDEX_URL, m[1].replace("&amp;", "&")), as_of


def _header(rows: list[tuple[Any, ...]]) -> tuple[int, dict[str, int]]:
    for i, row in enumerate(rows[:6]):
        first = _norm(row[0]) if row else ""
        if first.startswith("queue"):
            names = [_norm(c) for c in row]
            if first == "queue" and i + 1 < len(rows):  # two-row header ('Queue' over 'Pos.')
                names = [_norm(f"{a or ''} {b or ''}") for a, b in zip(row, rows[i + 1], strict=False)]
                i += 1
            cols = {}
            for fld, options in COLUMNS.items():
                idx = next((names.index(o) for o in options if o in names), None)
                if idx is not None:
                    cols[fld] = idx
            if "queue_id" not in cols:
                break
            return i + 1, cols
    raise ParseError("NYISO sheet without a 'Queue Pos.' header row")


def _qid(value: Any) -> str | None:
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    if isinstance(value, int):
        return f"{value:04d}"
    s = text(value)
    if s is None:
        return None
    s = s.upper()
    if s.isdigit():
        s = s.zfill(4)
    return s if _QID.match(s) else None


def _code(value: Any) -> str | None:
    s = text(value)
    if s is None:
        return None
    return s.upper().removesuffix("C") if s.upper() != "C" else s.upper()


def parse_sheet(name: str, rows: list[tuple[Any, ...]]) -> list[QueueEntry]:
    key, default_status = SHEETS[name]
    start, cols = _header(rows)

    def cell(row: tuple[Any, ...], fld: str) -> Any:
        i = cols.get(fld)
        return row[i] if i is not None and i < len(row) else None

    out = []
    for row in rows[start:]:
        qid = _qid(cell(row, "queue_id"))
        if qid is None:
            continue
        code = _code(cell(row, "status"))
        status = key.get(code, f"status {code}") if code else default_status
        fuel_code = (text(cell(row, "fuel")) or "").upper()
        fuel = FUEL.get(fuel_code, fuel_code or None)
        if name == "Load Projects":
            fuel = LOAD
        end_use = (text(cell(row, "end_use")) or "").upper()
        studies = text(cell(row, "studies"))
        out.append(QueueEntry(qid, {
            "name": text(cell(row, "name")),
            "developer": text(cell(row, "developer")),
            "transmission_owner": text(cell(row, "utility")),
            "county": text(cell(row, "county")),
            "state": text(cell(row, "state")),
            "poi": text(cell(row, "poi")),
            "fuel": fuel,
            "project_type": text(cell(row, "record_type")) or name,
            "mw_requested": number(cell(row, "mw")),
            "status": status,
            "queue_date": us_date(cell(row, "queue_date")),
            "proposed_cod": month_or_date(cell(row, "cod")) or month_or_date(cell(row, "in_service")),
            "study_phase": f"Studies available: {studies}" if studies and studies.upper() != "N/A" else None,
            "zone": text(cell(row, "zone")),
            "end_use": f"{end_use} ({END_USE[end_use]})" if end_use in END_USE else (end_use or None),
        }, updated=us_date(cell(row, "updated"))))
    return out


def parse_workbook(data: bytes) -> list[QueueEntry]:
    """Every queue entry in the workbook; an entry listed on several sheets keeps its first (active) row."""
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises zipfile/KeyError/InvalidFileException for non-workbooks
        raise ParseError(f"NYISO queue file is not an xlsx workbook: {type(exc).__name__}: {exc}") from None
    by_name = {" ".join(n.split()): n for n in wb.sheetnames}
    missing = [s for s in ("Interconnection Queue", "Load Projects", "Withdrawn") if s not in by_name]
    if missing:
        raise ParseError(f"NYISO workbook lacks sheets {missing}; has {wb.sheetnames}")
    seen: dict[str, QueueEntry] = {}
    for sheet in SHEETS:
        if sheet in by_name:
            for e in parse_sheet(sheet, list(wb[by_name[sheet]].iter_rows(values_only=True))):
                seen.setdefault(e.queue_id, e)
    return list(seen.values())


def load(index: DocumentRow, workbook: DocumentRow, as_of: date, stats: RunStats) -> None:
    entries = parse_workbook(workbook.read())
    if len(entries) < 500:
        raise ParseError(f"NYISO workbook looks truncated: {len(entries)} entries")
    with db.transaction() as conn:
        store = QueueStore(conn, "nyiso", "NYISO", PRODUCER, as_of, stats)
        for e in entries:
            store.write(e, source_url=workbook.url, document_id=workbook.id)
        store.finish(source_url=workbook.url, document_id=workbook.id)
    stats.extra["nyiso_workbook"] = {"url": workbook.url, "as_of": as_of.isoformat(), "index_document": index.id}


def run(*, fetcher: Fetcher | None = None) -> dict[str, Any]:
    def body(stats: RunStats) -> None:
        index = fetch_document(SOURCE_ID, INDEX_URL, doc_type="index_html", title="NYISO Interconnections",
                               fetcher=fetcher)
        stats.document(index)
        url, as_of = find_workbook(index.read())
        as_of = min(as_of, index.fetched_at.astimezone(ET).date())
        workbook = fetch_document(SOURCE_ID, url, doc_type="queue_xlsx", title=f"NYISO Interconnection Queue {as_of}",
                                  doc_date=as_of, external_id=as_of.isoformat(), fetcher=fetcher)
        stats.document(workbook)
        load(index, workbook, as_of, stats)

    return collect(SOURCE, body)


@app.periodic(cron=CRON_UTC, periodic_id="weekly_wed")
@app.task(name="collect.nyiso_queue", queue="collect", queueing_lock="collect.nyiso_queue",
          lock="collect.nyiso_queue")
def collect_task(timestamp: int | None = None) -> dict[str, Any]:
    return run()


@cli.action("collect", "nyiso_queue", help="NYISO interconnection queue workbook (incl. large loads)")
def collect_cli(args: argparse.Namespace) -> int:
    return print_stats(run())
