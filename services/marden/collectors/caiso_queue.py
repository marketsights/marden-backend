"""CAISO generator interconnection queue -> queue_project facts (tag F, source CAISO). See iso_queue.

CAISO's Library page "Public queue report" (https://www.caiso.com/library/public-queue-report) posts the
report as https://www.caiso.com/documents/publicqueuereport.xlsx (and a PDF of the same), regenerated
daily ("Report Run Date: MM/DD/YYYY" in each sheet's first row). The workbook is stored raw and parsed
in code. Sheets: "Grid GenerationQueue" (active), "Completed Generation Projects" (the legend: "Complete
= project is in Commercial Operation") and "Withdrawn Generation Projects". Subjects are
`caiso:<Queue Position>` (22, 643R, 1A ...).

Fields: name = "Project Name"; transmission_owner = "Utility" (PGAE, SCE, SDGE, GLW, VEA, DCRT); county,
state; poi = "Station or Transmission Line"; fuel = "Fuel-1..3" joined; project_type = "Type-1..3" with
the deliverability ("Full Capacity, Partial or Energy Only"); mw_requested = "Net MWs to Grid" (else
"MW-1"); status = "Application Status" ('ACTIVE (Suspended)' when "Suspension Status" says so);
queue_date = "Queue Date"; proposed_cod = "Current On-line Date"; actual_cod = "Actual On-line Date"
(completed); withdrawn_date = "Withdrawn Date"; study_phase = the published study and agreement columns
(and "Reason for Withdrawal"); cycle = "Study Process" (C14, Serial LGIP ...); zone = "PTO Study Region".
CAISO publishes no interconnection customer (developer) and no load requests: large loads connect
through the utilities (PG&E, SCE, SDG&E), not the CAISO queue. Matching to our sites therefore relies on
the project name alone (queue-match@0.1, county + state + name).

Cadence: regenerated daily; collected weekly (Friday). `python -m marden collect caiso_queue`.
"""

from __future__ import annotations

import argparse
import io
import re
from datetime import date
from typing import Any

import openpyxl

from marden import cli
from marden.collectors.iso_queue import QueueEntry, QueueStore, month_or_date, number, text, us_date
from marden.collectors.rates_common import ET, ParseError, RunStats, collect, print_stats
from marden.core import db, http
from marden.core.documents import DocumentRow
from marden.core.http import Fetcher, fetch_document
from marden.core.queue import app
from marden.core.sources import declare_source

SOURCE_ID = "caiso_queue"
SOURCE_NAME = "CAISO"
PRODUCER = "collector:caiso_queue"
HOST = "www.caiso.com"
PAGE_URL = f"https://{HOST}/library/public-queue-report"
XLSX_URL = f"https://{HOST}/documents/publicqueuereport.xlsx"
CRON_UTC = "41 12 * * 5"  # Fridays 12:41 UTC (the report is regenerated around 08:00 UTC)
TERMS_URL = f"https://{HOST}/privacy-terms-of-use"
TERMS_NOTE = (
    f"Checked 2026-10-21. California ISO Terms of Use ({TERMS_URL}, last updated 10/02/2018), Restrictions on "
    "Use of Materials: \"Most of the materials and information contained on this Website ... were generated, "
    "compiled or assembled from materials and information that are freely available for public use consistent "
    "with the general policies of the Public Records Act (California Government Code section 6250, et seq.), "
    "as provided by California Public Utilities Code section 345.5, and may be used by you provided that you "
    "keep intact all copyright, trademark and other proprietary notices and that you credit the California ISO "
    "when using such materials and/or information.\" Every fact carries source_name CAISO and the report URL. "
    "robots.txt (https://www.caiso.com/robots.txt): \"Disallow: /resources/*\" only; /documents/ and /library/ "
    "are allowed. The workbook is the public download on the Library page (no login; CAISO's separate API "
    "Terms of Use apply to its APIs, which are not used). One request a week."
)
SOURCE = declare_source(SOURCE_ID, "California ISO", "portal", f"https://{HOST}", 7 * 24 * 60, TERMS_NOTE)
http.set_rate_limit(HOST, 0.5)

SHEETS = {  # sheet -> status class of its rows
    "Grid GenerationQueue": "active",
    "Completed Generation Projects": "in_service",
    "Withdrawn Generation Projects": "withdrawn",
}
COLUMNS = {  # field -> published header names (whitespace collapsed)
    "name": ("Project Name", "Project Name - Confidential"),
    "queue_id": ("Queue Position",),
    "received": ("Interconnection Request Receive Date",),
    "queue_date": ("Queue Date",),
    "status": ("Application Status",),
    "withdrawn": ("Withdrawn Date",),
    "process": ("Study Process",),
    "net_mw": ("Net MWs to Grid",),
    "mw1": ("MW-1",),
    "deliverability": ("Full Capacity, Partial or Energy Only (FC/P/EO)",),
    "county": ("County",),
    "state": ("State",),
    "utility": ("Utility",),
    "region": ("PTO Study Region",),
    "poi": ("Station or Transmission Line",),
    "current_cod": ("Current On-line Date",),
    "actual_cod": ("Actual On-line Date",),
    "filed_cod": ("Proposed On-line Date (as filed with IR)",),
    "suspension": ("Suspension Status",),
    "reason": ("Reason for Withdrawal",),
}
STUDIES = (  # published study / agreement columns -> label in study_phase
    ("Feasibility Study or Supplemental Review", "Feasibility/Supplemental Review"),
    ("System Impact Study or Phase I Cluster Study", "SIS/Phase I"),
    ("Facilities Study (FAS) or Phase II Cluster Study", "FAS/Phase II"),
    ("Optional Study (OS)", "Optional Study"),
    ("Interconnection Agreement Status", "IA"),
)
_RUN_DATE = re.compile(r"Report Run Date:\s*(\d{1,2}/\d{1,2}/\d{4})")
MIN_ENTRIES = 1000  # CAISO lists ~2,300 projects; fewer means a truncated or changed workbook


def _norm(value: Any) -> str:
    return " ".join(str(value).split()) if value is not None else ""


def _joined(row: tuple[Any, ...], header: list[str], prefix: str) -> str | None:
    parts = [text(row[i]) for i, h in enumerate(header) if re.fullmatch(rf"{prefix}-\d", h) and i < len(row)]
    return " + ".join(p for p in parts if p) or None


def _qid(value: Any) -> str | None:
    s = text(value)
    return s.upper() if s and re.fullmatch(r"[0-9]+[A-Za-z]{0,3}", s) else None


def parse_sheet(name: str, rows: list[tuple[Any, ...]]) -> tuple[date | None, list[QueueEntry]]:
    cls = SHEETS[name]
    run_date = None
    for row in rows[:3]:
        for c in row:
            m = _RUN_DATE.search(_norm(c))
            if m:
                run_date = us_date(m[1])
    hdr_i = next((i for i, r in enumerate(rows[:8]) if "Queue Position" in [_norm(c) for c in r]), None)
    if hdr_i is None:
        raise ParseError(f"CAISO sheet {name!r} has no 'Queue Position' header")
    header = [_norm(c) for c in rows[hdr_i]]
    cols = {f: next((header.index(o) for o in opts if o in header), None) for f, opts in COLUMNS.items()}
    missing = [f for f in ("name", "queue_id", "status", "county", "state") if cols[f] is None]
    if missing:
        raise ParseError(f"CAISO sheet {name!r} lacks columns {missing}: {header}")
    studies = [(header.index(h), label) for h, label in STUDIES if h in header]

    def cell(row: tuple[Any, ...], fld: str) -> Any:
        i = cols[fld]
        return row[i] if i is not None and i < len(row) else None

    out = []
    for row in rows[hdr_i + 1:]:
        qid = _qid(cell(row, "queue_id"))
        if qid is None:
            continue  # legend and footnote rows
        status = text(cell(row, "status"))
        entry_cls = cls
        if cls == "active" and (text(cell(row, "suspension")) or "").lower() == "suspended":
            status, entry_cls = f"{status or 'ACTIVE'} (Suspended)", "suspended"
        phases = [f"{label}: {v}" for i, label in studies if i < len(row) and (v := text(row[i]))]
        if reason := text(cell(row, "reason")):
            phases.append(f"Reason for withdrawal: {reason}")
        types, deliver = _joined(row, header, "Type"), text(cell(row, "deliverability"))
        cod = month_or_date(cell(row, "current_cod")) or month_or_date(cell(row, "filed_cod"))
        out.append(QueueEntry(qid, {
            "name": text(cell(row, "name")),
            "transmission_owner": text(cell(row, "utility")),
            "county": text(cell(row, "county")),
            "state": (text(cell(row, "state")) or "").upper() or None,
            "poi": text(cell(row, "poi")),
            "fuel": _joined(row, header, "Fuel"),
            "project_type": f"{types} ({deliver})" if types and deliver else types,
            "mw_requested": number(cell(row, "net_mw")) or number(cell(row, "mw1")),
            "status": status,
            "status_class": entry_cls,
            "queue_date": us_date(cell(row, "queue_date")) or us_date(cell(row, "received")),
            "proposed_cod": None if cls == "in_service" else cod,
            "actual_cod": us_date(cell(row, "actual_cod")) if cls == "in_service" else None,
            "withdrawn_date": us_date(cell(row, "withdrawn")),
            "study_phase": "; ".join(phases) or None,
            "cycle": text(cell(row, "process")),
            "zone": text(cell(row, "region")),
        }))
    return run_date, out


def parse_workbook(data: bytes) -> tuple[date | None, list[QueueEntry]]:
    """(the report run date, every entry). A queue position on two sheets keeps its first (active) row."""
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises zipfile/KeyError/InvalidFileException for non-workbooks
        raise ParseError(f"CAISO queue file is not an xlsx workbook: {type(exc).__name__}: {exc}") from None
    by_name = {_norm(n): n for n in wb.sheetnames}
    missing = [s for s in SHEETS if s not in by_name]
    if missing:
        raise ParseError(f"CAISO workbook lacks sheets {missing}; has {wb.sheetnames}")
    seen: dict[str, QueueEntry] = {}
    run_dates: set[date] = set()
    for sheet in SHEETS:
        run_date, entries = parse_sheet(sheet, list(wb[by_name[sheet]].iter_rows(values_only=True)))
        if run_date:
            run_dates.add(run_date)
        for e in entries:
            seen.setdefault(e.queue_id, e)
    return (max(run_dates) if run_dates else None), list(seen.values())


def load(doc: DocumentRow, stats: RunStats, *, min_entries: int = MIN_ENTRIES) -> None:
    run_date, entries = parse_workbook(doc.read())
    if len(entries) < min_entries:
        raise ParseError(f"CAISO queue workbook looks truncated: {len(entries)} entries")
    fetched = doc.fetched_at.astimezone(ET).date()
    day = min(run_date or fetched, fetched)
    with db.transaction() as conn:
        store = QueueStore(conn, "caiso", SOURCE_NAME, PRODUCER, day, stats)
        for e in entries:
            store.write(e, source_url=XLSX_URL, document_id=doc.id)
        store.finish(source_url=XLSX_URL, document_id=doc.id)
    stats.extra["caiso_report"] = {"url": XLSX_URL, "page": PAGE_URL,
                                   "run_date": run_date.isoformat() if run_date else None}


def run(*, fetcher: Fetcher | None = None) -> dict[str, Any]:
    def body(stats: RunStats) -> None:
        doc = fetch_document(SOURCE_ID, XLSX_URL, doc_type="queue_xlsx", title="CAISO public queue report",
                             fetcher=fetcher)
        stats.document(doc)
        load(doc, stats)

    return collect(SOURCE, body)


@app.periodic(cron=CRON_UTC, periodic_id="weekly_fri")
@app.task(name="collect.caiso_queue", queue="collect", queueing_lock="collect.caiso_queue",
          lock="collect.caiso_queue")
def collect_task(timestamp: int | None = None) -> dict[str, Any]:
    return run()


@cli.action("collect", "caiso_queue", help="CAISO generator interconnection queue (public queue report xlsx)")
def collect_cli(args: argparse.Namespace) -> int:
    return print_stats(run())
