"""LBNL "Queued Up" national interconnection queue dataset -> queue_project facts (tag R). See iso_queue.

Berkeley Lab and GridTracker compile, clean and publish the generation interconnection queues of the 7 US
ISO/RTOs and ~50 non-ISO transmission providers (Southern, Duke, TVA, Santee Cooper, Dominion Energy SC,
BPA, PacifiCorp, APS, SRP, TEP, NV Energy, PNM, El Paso Electric, PSCo, Idaho Power, PGE, PSE ...) once a
year ("Queued Up: <year> Edition", data through the end of the prior year). It is a third-party compilation
of filed queue data, so every fact is tag R with source_name 'LBNL Queued Up', valid_from = the data's
as-of (Dec 31 of the year in the file name). It excludes load (large-load) requests.

- Discovery: the series page https://eta-publications.lbl.gov/research-areas/queued-characteristics-power-plants
  lists every edition's data file; the newest lbnl_ix_queue_data_file_thru<year>[_v<n>].xlsx is used. Its
  edition page is fetched too and must still say the data file is licensed CC BY 4.0 (otherwise nothing is
  written). emp.lbl.gov (the "latest edition" landing page) answers Cloudflare "Sorry, you have been
  blocked" and is not used.
- Subjects are separate from the operators' own entries: 'lbnl:<entity>:<q_id>' (entity = LBNL's
  transmission-provider code, lower case: 'lbnl:soco:IC-525', 'lbnl:pjm:AE1-070'), so LBNL's year-end
  values never displace an operator's current ones in fact_current. When the operator's entry is in
  Marden (pjm, nyiso, isone, bpa, miso, ercot, spp, caiso queues), `operator_entry` names it: the cross-check.
- Fields: name, iso (the ISO/RTO, none for non-ISO areas), transmission_owner (LBNL 'utility'), developer,
  county, state, poi, fuel (type_clean), project_type, service, mw_requested (mw_1, the capacity of the
  first resource type; hybrids also get mw_by_type), status (q_status), queue_date, proposed_cod,
  actual_cod, withdrawn_date, ia_date, study_phase (IA_phase_clean), cycle (cluster).
- Matching to our sites/campuses/plants, link facts and events are the shared queue-match rules.

Cadence: LBNL publishes yearly (spring); checked monthly (9th). `python -m marden collect lbnl_queued_up`.
"""

from __future__ import annotations

import argparse
import io
import re
from collections import Counter
from datetime import date, datetime
from typing import Any
from urllib.parse import urljoin

import openpyxl

from marden import cli
from marden.collectors.iso_queue import QueueEntry, QueueStore, number, text, us_date
from marden.collectors.rates_common import ParseError, RunStats, collect, print_stats
from marden.core import db, http
from marden.core.db import Conn
from marden.core.documents import DocumentRow
from marden.core.http import Fetcher, fetch_document
from marden.core.queue import app
from marden.core.sources import declare_source

SOURCE_ID = "lbnl_queued_up"
SOURCE_NAME = "LBNL Queued Up"
PRODUCER = f"collector:{SOURCE_ID}"
HOST = "eta-publications.lbl.gov"
SERIES_URL = f"https://{HOST}/research-areas/queued-characteristics-power-plants"
CRON_UTC = "7 12 9 * *"  # the 9th of each month, 12:07 UTC
ATTRIBUTION = (
    "Lawrence Berkeley National Laboratory and GridTracker. Queued Up Data File. Copyright (c) 2026, The Regents "
    "of the University of California, through Lawrence Berkeley National Laboratory (subject to receipt of any "
    "required approvals from the U.S. Dept. of Energy) & GridTracker. CC BY 4.0."
)
LICENCE_TEXT = "The Queued Up data file is licensed CC BY 4.0"
TERMS_NOTE = (
    "Checked 2026-10-21. Edition page (https://eta-publications.lbl.gov/publications/queued-2026-edition-"
    "characteristics): \"Copyright Notice: The Queued Up data file is licensed CC BY 4.0. You may use, share, or "
    "adapt the dataset as long as you attribute it to Lawrence Berkeley National Laboratory and GridTracker. "
    f"{ATTRIBUTION}\" Commercial use is allowed under CC BY 4.0 with that attribution (shown with every LBNL "
    "value; licence entry 'LBNL Queued Up' in marden.core.licensing). robots.txt "
    "(https://eta-publications.lbl.gov/robots.txt) disallows only /core/, /profiles/, admin, search and user "
    "paths; /research-areas/, /publications/ and /sites/default/files/ are allowed, no crawl-delay. "
    "emp.lbl.gov (robots.txt Crawl-delay: 90) answers Cloudflare 'Sorry, you have been blocked' and is not used. "
    "Each run re-checks the licence sentence on the edition page; 3 requests a month (the 15 MB file only "
    "when it changed, by conditional GET)."
)
SOURCE = declare_source(SOURCE_ID, SOURCE_NAME, "feed", f"https://{HOST}", 31 * 24 * 60, TERMS_NOTE)
http.set_rate_limit(HOST, 0.2)

_FILE = re.compile(r'href="([^"]*?/lbnl_ix_queue_data_file_thru(\d{4})(?:_v(\d+))?\.xlsx)"', re.I)
_EDITION = re.compile(r'href="((?:https://eta-publications\.lbl\.gov)?/publications/queued-(\d{4})-edition[^"]*)"',
                      re.I)
SHEET = "03. Complete Queue Data"
HEADER = ("q_id", "q_status", "q_date", "prop_date", "on_date", "wd_date", "ia_date", "IA_phase_raw",
          "IA_phase_clean", "county", "state", "fips_code", "poi_name", "region", "project_name", "utility",
          "entity", "developer", "cluster", "service", "project_type", "type_1", "type_2", "type_3", "type_clean",
          "mw_1", "mw_2", "mw_3", "q_year", "prop_year")
ISOS = frozenset({"PJM", "MISO", "ERCOT", "CAISO", "SPP", "NYISO", "ISO-NE"})
STATUS_CLASS = {"active": "active", "withdrawn": "withdrawn", "suspended": "suspended", "operational": "in_service"}
# LBNL entity -> Marden's queue_project prefix for operators Marden collects directly.
OPERATORS = {"PJM": "pjm", "NYISO": "nyiso", "ISO-NE": "isone", "BPA": "bpa", "MISO": "miso", "ERCOT": "ercot",
             "SPP": "spp", "CAISO": "caiso"}
MIN_ENTRIES = 20_000  # the 2026 edition has 38,201 rows; fewer means a truncated or changed file


def find_data_file(page: bytes) -> tuple[str, int]:
    """The newest data file on the series page and the year its data runs through."""
    found = [(int(m[2]), int(m[3] or 0), m[1]) for m in _FILE.finditer(page.decode("utf-8", errors="replace"))]
    if not found:
        raise ParseError(f"no lbnl_ix_queue_data_file_thru<year>.xlsx link on {SERIES_URL}")
    year, _, href = max(found)
    return urljoin(SERIES_URL, href), year


def find_edition_page(page: bytes, data_year: int) -> str:
    """The '<data year + 1> Edition' publication page."""
    for m in _EDITION.finditer(page.decode("utf-8", errors="replace")):
        if int(m[2]) == data_year + 1:
            return urljoin(SERIES_URL, m[1])
    raise ParseError(f"no Queued Up {data_year + 1} Edition page on {SERIES_URL}")


def check_licence(edition: bytes, data_url: str) -> None:
    body = edition.decode("utf-8", errors="replace")
    if LICENCE_TEXT not in " ".join(re.sub(r"<[^>]+>", " ", body).split()):
        raise ParseError("the edition page no longer says the data file is licensed CC BY 4.0; nothing written")
    if data_url.rsplit("/", 1)[-1] not in body:
        raise ParseError(f"the edition page does not link {data_url}")


def _qid(value: Any) -> str | None:
    if isinstance(value, datetime):
        return value.date().isoformat()  # an id Excel turned into a date
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return text(value)


def _slug(entity: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", entity.lower()).strip("-")


def _values(r: dict[str, Any]) -> dict[str, Any]:
    status = (text(r["q_status"]) or "").lower() or None
    parts = [(text(r[f"type_{n}"]), number(r[f"mw_{n}"])) for n in (1, 2, 3)]
    parts = [(t, mw) for t, mw in parts if t and mw is not None]
    by_type: dict[str, int | float] | None = None
    if len(parts) > 1:
        by_type = {}
        for t, mw in parts:
            by_type[t] = by_type.get(t, 0) + mw
    region = text(r["region"])
    return {
        "name": text(r["project_name"]),
        "iso": region if region in ISOS else None,
        "transmission_owner": text(r["utility"]),
        "developer": text(r["developer"]),
        "county": text(r["county"]),
        "state": text(r["state"]),
        "poi": text(r["poi_name"]),
        "fuel": text(r["type_clean"]),
        "project_type": text(r["project_type"]),
        "service": text(r["service"]),
        "mw_requested": number(r["mw_1"]),
        "mw_by_type": by_type,
        "status": status,
        "status_class": STATUS_CLASS.get(status) if status else None,
        "queue_date": us_date(r["q_date"]),
        "proposed_cod": d.isoformat() if (d := us_date(r["prop_date"])) else None,
        "actual_cod": us_date(r["on_date"]),
        "withdrawn_date": us_date(r["wd_date"]),
        "ia_date": us_date(r["ia_date"]),
        "study_phase": text(r["IA_phase_clean"]),
        "cycle": text(r["cluster"]),
    }


def parse_workbook(data: bytes) -> list[tuple[str, str, QueueEntry]]:
    """[(entity, the entity's q_id as published, entry)] with entry ids '<entity slug>:<q_id>'. A q_id
    listed twice for one entity gets '/<q_date>' (and '-2', '-3' if still shared) so each row is kept."""
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises zipfile/KeyError/InvalidFileException for non-workbooks
        raise ParseError(f"LBNL file is not an xlsx workbook: {type(exc).__name__}: {exc}") from None
    if SHEET not in wb.sheetnames:
        raise ParseError(f"LBNL workbook lacks sheet {SHEET!r}; has {wb.sheetnames[:6]}")
    rows = wb[SHEET].iter_rows(values_only=True)
    header: tuple[Any, ...] | None = None
    for row in rows:
        if row and row[0] == "q_id":
            header = row
            break
    if header is None or tuple(header[:len(HEADER)]) != HEADER:
        raise ParseError(f"LBNL queue columns changed: {header}")
    raw: list[tuple[str, str, dict[str, Any]]] = []
    for row in rows:
        r = dict(zip(HEADER, row, strict=False))
        entity, qid = text(r["entity"]), _qid(r["q_id"])
        if entity and qid:
            raw.append((entity, qid, r))
    counts = Counter((e, q) for e, q, _ in raw)
    used: set[str] = set()
    out = []
    for entity, qid, r in raw:
        base = f"{_slug(entity)}:{qid}"
        if counts[(entity, qid)] > 1:
            day = us_date(r["q_date"])
            base = f"{base}/{day.isoformat() if day else 'undated'}"
        eid, n = base, 2
        while eid in used:
            eid, n = f"{base}-{n}", n + 1
        used.add(eid)
        out.append((entity, qid, QueueEntry(eid, _values(r))))
    return out


def operator_ids(conn: Conn) -> set[str]:
    rows = conn.execute("SELECT id FROM queue_project WHERE iso = ANY(%s)", (sorted(set(OPERATORS.values())),))
    return {r["id"] for r in rows.fetchall()}


def operator_entry(entity: str, qid: str, known: set[str]) -> str | None:
    prefix = OPERATORS.get(entity)
    if prefix is None:
        return None
    candidates = [qid, qid.zfill(4)] if prefix == "nyiso" and qid.isdigit() else [qid]
    return next((f"{prefix}:{c}" for c in candidates if f"{prefix}:{c}" in known), None)


def load(doc: DocumentRow, data_year: int, stats: RunStats, *, min_entries: int = MIN_ENTRIES) -> None:
    entries = parse_workbook(doc.read())
    if len(entries) < min_entries:
        raise ParseError(f"LBNL queue file looks truncated: {len(entries)} rows")
    as_of = date(data_year, 12, 31)
    with db.transaction() as conn:
        known = operator_ids(conn)
        store = QueueStore(conn, "lbnl", SOURCE_NAME, PRODUCER, as_of, stats, tag="R", iso_label="")
        cross = Counter()
        for entity, qid, e in entries:
            if op := operator_entry(entity, qid, known):
                e.values["operator_entry"] = op
                cross[entity] += 1
            store.write(e, source_url=doc.url, document_id=doc.id)
        store.finish(source_url=doc.url, document_id=doc.id)
    stats.extra["lbnl"] = {
        "data_through": as_of.isoformat(), "rows": len(entries),
        "by_entity": dict(Counter(e for e, _, _ in entries).most_common()),
        "operator_entry_matches": dict(sorted(cross.items())), "attribution": ATTRIBUTION,
    }


def run(*, fetcher: Fetcher | None = None) -> dict[str, Any]:
    def body(stats: RunStats) -> None:
        series = fetch_document(SOURCE_ID, SERIES_URL, doc_type="index_html", title="LBNL Queued Up editions",
                                fetcher=fetcher)
        stats.document(series)
        data_url, year = find_data_file(series.read())
        edition_url = find_edition_page(series.read(), year)
        edition = fetch_document(SOURCE_ID, edition_url, doc_type="index_html",
                                 title=f"Queued Up {year + 1} Edition", fetcher=fetcher)
        stats.document(edition)
        check_licence(edition.read(), data_url)
        doc = fetch_document(SOURCE_ID, data_url, doc_type="queue_xlsx", title=f"LBNL queue data through {year}",
                             doc_date=date(year, 12, 31), external_id=str(year), fetcher=fetcher)
        stats.document(doc)
        load(doc, year, stats)

    return collect(SOURCE, body)


@app.periodic(cron=CRON_UTC, periodic_id="monthly_9th")
@app.task(name="collect.lbnl_queued_up", queue="collect", queueing_lock="collect.lbnl_queued_up",
          lock="collect.lbnl_queued_up")
def collect_task(timestamp: int | None = None) -> dict[str, Any]:
    return run()


@cli.action("collect", "lbnl_queued_up", help="LBNL Queued Up national interconnection queue dataset (tag R)")
def collect_cli(args: argparse.Namespace) -> int:
    return print_stats(run())
