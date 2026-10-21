"""EIA-930 reference tables -> every balancing authority EIA-930 reports, as `grid_region` rows 'ba:<code>'.

Source: https://www.eia.gov/electricity/930-content/EIA930_Reference_Tables.xlsx (sheet "BAs": BA Code, BA Name,
Time Zone, Region/Country Code, Region/Country Name, Generation Only BA, Demand by BA Subregion, U.S. BA,
Active BA, Activation Date, Retirement Date). Retired BAs stay as rows with `active = false` and their
retirement date, so the spatial join can flag a boundary drawn before a BA was retired.

Interconnection (the row's parent) follows the EIA-930 region of the BA: TEX is the ERCOT Interconnection;
CAL, NW and SW are the Western Interconnection; every other US region is the Eastern Interconnection;
Canadian BAs follow grid_registry.CANADA; Mexico's CEN/CFE get no parent.

TERMS (checked 2026-10-21): see TERMS_NOTE. Cadence: monthly (EIA revises the table when BAs change, e.g. SWPW
and BHBA added and WACM/WAUW retired on 2026-04-01). `python -m marden collect eia930_ref`.
"""

from __future__ import annotations

import argparse
import io
from datetime import date, datetime
from email.utils import parsedate_to_datetime
from typing import Any

from marden import cli
from marden.collectors import grid_registry
from marden.collectors.grid_store import Region, UpsertStats, upsert
from marden.collectors.rates_common import ParseError, RunStats, collect, print_stats
from marden.core import db
from marden.core.documents import DocumentRow
from marden.core.http import Fetcher, fetch_document
from marden.core.queue import app
from marden.core.sources import SourceSpec, declare_source

SOURCE_ID = "eia930_ref"
URL = "https://www.eia.gov/electricity/930-content/EIA930_Reference_Tables.xlsx"
REUSE_URL = "https://www.eia.gov/about/copyrights_reuse.php"
CRON_UTC = "13 11 2 * *"  # 2nd of each month, 11:13 UTC
TERMS_NOTE = (
    "Checked 2026-10-21. EIA Copyrights and Reuse (" + REUSE_URL + "): \"U.S. government publications are in the "
    "public domain and are not subject to copyright protection. You may use and/or distribute any of our data, "
    "files, databases, reports, graphs, charts, and other information products that are on our website or that "
    "you receive through our email distribution service. However, if you use or reproduce any of our information "
    "products, you should use an acknowledgment, which includes the publication date, such as: \"Source: U.S. "
    "Energy Information Administration (Oct 2008).\"\" robots.txt (https://www.eia.gov/robots.txt): \"User-agent: * "
    "Allow: /\"; /electricity/ is not disallowed. One file per run, conditional GET, 2 req/s host limit."
)
SOURCE = declare_source(SOURCE_ID, "EIA-930 reference tables", "feed", "https://www.eia.gov/electricity/gridmonitor/",
                        30 * 24 * 60, TERMS_NOTE)

HEADER = ("BA Code", "BA Name", "Time Zone", "Region/Country Code", "Region/Country Name", "Generation Only BA",
          "Demand by BA Subregion", "U.S. BA", "Active BA", "Activation Date", "Retirement Date")
WESTERN_REGIONS = frozenset({"CAL", "NW", "SW"})
ERCOT_REGIONS = frozenset({"TEX"})
EASTERN_REGIONS = frozenset({"CAR", "CENT", "FLA", "MIDA", "MIDW", "NE", "NY", "SE", "TEN"})
COUNTRY_BY_REGION = {"CAN": "CA", "MEX": "MX"}
CANADA_PARENT = {code: parent for _, _, _, parent, code, _ in grid_registry.CANADA if code}
# EIA-930 files the retired WAPA Upper Great Plains East BA (in SPP since 2015-10-01) under region NW, but it
# was in the Eastern Interconnection (UGP West, WAUW, is the Western part).
EXCEPTIONS = {"WAUE": grid_registry.EASTERN}


def _yes(value: Any, column: str, code: str) -> bool:
    text = str(value or "").strip().lower()
    if text in ("yes", "y"):
        return True
    if text in ("no", "n"):
        return False
    raise ParseError(f"EIA-930 BA {code}: {column} is {value!r}, expected Yes or No")


def _day(value: Any) -> str | None:
    if value is None or str(value).strip() == "":
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    try:
        return datetime.strptime(str(value).strip(), "%m/%d/%Y").date().isoformat()
    except ValueError:
        raise ParseError(f"EIA-930 reference table: unparseable date {value!r}") from None


def parse(data: bytes) -> list[dict[str, Any]]:
    """The 'BAs' sheet as dicts (one per BA code, in file order)."""
    import openpyxl

    try:
        book = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises several types for a file that is not a workbook
        raise ParseError(f"EIA-930 reference tables are not a readable workbook: {exc}") from None
    if "BAs" not in book.sheetnames:
        raise ParseError(f"EIA-930 reference tables have no 'BAs' sheet (sheets: {book.sheetnames})")
    rows = book["BAs"].iter_rows(values_only=True)
    header = tuple(str(h).strip() if h is not None else "" for h in next(rows, ()))
    if header[: len(HEADER)] != HEADER:
        raise ParseError(f"EIA-930 'BAs' sheet: unexpected header {header!r}")
    out: list[dict[str, Any]] = []
    seen: set[str] = set()
    for values in rows:
        if not values or not values[0] or not str(values[0]).strip():
            continue
        cells = dict(zip(HEADER, values, strict=False))
        code = str(cells["BA Code"]).strip()
        if code in seen:
            raise ParseError(f"EIA-930 'BAs' sheet lists {code} twice")
        seen.add(code)
        out.append({
            "code": code,
            "name": str(cells["BA Name"] or "").strip(),
            "time_zone": str(cells["Time Zone"] or "").strip() or None,
            "region_code": str(cells["Region/Country Code"] or "").strip(),
            "region_name": str(cells["Region/Country Name"] or "").strip(),
            "generation_only": _yes(cells["Generation Only BA"], "Generation Only BA", code),
            "demand_by_subregion": _yes(cells["Demand by BA Subregion"], "Demand by BA Subregion", code),
            "us_ba": _yes(cells["U.S. BA"], "U.S. BA", code),
            "active": _yes(cells["Active BA"], "Active BA", code),
            "activation_date": _day(cells["Activation Date"]),
            "retirement_date": _day(cells["Retirement Date"]),
        })
    if not out:
        raise ParseError("EIA-930 'BAs' sheet has no rows")
    return out


def interconnection_for(code: str, region_code: str) -> str | None:
    if code in EXCEPTIONS:
        return EXCEPTIONS[code]
    if region_code in ERCOT_REGIONS:
        return grid_registry.ERCOT
    if region_code in WESTERN_REGIONS:
        return grid_registry.WESTERN
    if region_code in EASTERN_REGIONS:
        return grid_registry.EASTERN
    if region_code == "CAN":
        return CANADA_PARENT.get(code)
    return None


def document_as_of(doc: DocumentRow) -> date:
    """The file's Last-Modified date (EIA does not date the workbook), else the day it was fetched."""
    modified = (doc.meta.get("http") or {}).get("last_modified")
    if modified:
        try:
            return parsedate_to_datetime(modified).date()
        except (TypeError, ValueError):
            pass
    return doc.fetched_at.date()


def region_for(ba: dict[str, Any], doc: DocumentRow) -> Region:
    meta = {**ba, "document_id": doc.id, "source_url": doc.url, "as_of": document_as_of(doc).isoformat()}
    country = COUNTRY_BY_REGION.get(ba["region_code"], "US")
    return Region(f"ba:{ba['code']}", "balancing_authority", ba["name"] or ba["code"], country,
                  interconnection_for(ba["code"], ba["region_code"]), {"eia930": meta})


def load(doc: DocumentRow, stats: RunStats) -> dict[str, int]:
    bas = parse(doc.read())
    registry_owned = {r.id for r in grid_registry.regions()}
    out = UpsertStats()
    unmapped: list[str] = []
    with db.transaction() as conn:
        grid_registry.ensure_registry(conn)
        for ba in bas:
            region = region_for(ba, doc)
            if region.parent_id is None and ba["region_code"] != "MEX":
                unmapped.append(f"{ba['code']} ({ba['region_code']})")
            out.add(upsert(conn, region, prefer_existing=region.id in registry_owned))
    stats.facts_seen += len(bas)
    stats.facts_new += out.inserted + out.changed
    stats.extra.update(bas=len(bas), active=sum(b["active"] for b in bas), us=sum(b["us_ba"] for b in bas),
                       rows=out.as_dict(), no_interconnection=unmapped)
    return out.as_dict()


def run(*, source: SourceSpec = SOURCE, fetcher: Fetcher | None = None) -> dict[str, Any]:
    def body(stats: RunStats) -> None:
        doc = fetch_document(source.id, URL, doc_type="eia930_reference_xlsx", title="EIA-930 reference tables",
                             external_id="EIA930_Reference_Tables", fetcher=fetcher)
        stats.document(doc)
        load(doc, stats)

    return collect(source, body)


@app.periodic(cron=CRON_UTC, periodic_id="monthly")
@app.task(name="collect.eia930_ref", queue="collect", queueing_lock="collect.eia930_ref", lock="collect.grid")
def collect_task(timestamp: int | None = None) -> dict[str, Any]:
    return run()


@cli.action("collect", "eia930_ref", help="EIA-930 balancing authority list -> grid_region rows")
def collect_cli(args: argparse.Namespace) -> int:
    return print_stats(run())
