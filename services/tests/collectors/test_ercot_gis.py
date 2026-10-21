"""ercot_gis: the GIS workbook parser, the MIS listing (real captured page), the run (listing, one download per
report, twelve-month backfill oldest first, idempotent rerun, a new month loaded once), the stop rules (403,
429, challenge page: nothing retried or worked around), the robots/owner-decision terms, the schedule, and a
live check of the listing (network).

Captured files (data/ercot_gis): `mis_listing_15933_20261021.html.gz`, the MIS listing as served on 2026-10-21;
`GIS_Report_*_trimmed.xlsx`, ERCOT's August and September 2026 reports cut to their header rows and a few dozen
request rows (values as published, re-saved by openpyxl to stay under 200 KB; see test_real_reports). The
in-test workbook below follows the same layout for the flow tests, which need a distinct file per posting.
"""

from __future__ import annotations

import gzip
import io
import itertools
import re
from collections.abc import Iterator
from datetime import date, datetime
from pathlib import Path

import httpx
import openpyxl
import pytest

from marden import cli
from marden.collectors import ercot_gis as ercot
from marden.collectors import ercot_gis_parse as gis
from marden.collectors.rates_common import ParseError
from marden.core import documents, http, licensing, sources
from marden.core.queue import app

DATA = Path(__file__).parent / "data" / "ercot_gis"
INCAPSULA = Path(__file__).parent / "data" / "ercot" / "incapsula_403.html"
LISTING = gzip.decompress((DATA / "mis_listing_15933_20261021.html.gz").read_bytes())
SEP_2026 = "1281327706"  # doclookupId of GIS_Report_September2026 on the captured listing
LARGE_HEADER = ["INR", "Project Name", "GIM Study Phase", "Interconnecting Entity", "POI Location", "County",
                "CDR Reporting Zone", "Projected COD", "Fuel", "Technology", "Capacity (MW)", None, None,
                "Screening Study Started", "Screening Study Complete", "FIS Requested", "FIS Approved",
                "Economic Study Required", "IA Signed", None, "Air Permit", "GHG Permit", "Water Availability",
                "Meets Planning", "Meets All Planning", None, "Construction Start", "Construction End",
                "Approved for Energization", "Approved for Synchronization"]


def gis_workbook(mw: float = 360.5) -> bytes:
    wb = openpyxl.Workbook()
    wb.active.title = "Disclaimer and References"
    large = wb.create_sheet(gis.LARGE)
    for _ in range(3):
        large.append(["GIM Project Details - Large Generators"])
    large.append(LARGE_HEADER)
    large.append([None] * 11 + ["Change indicators: Proj Name, MW", "Approval Date for"])  # sub-header: no INR
    large.append(["24INR0459", "Test Peaker", "SS Completed, FIS Completed, IA", "Example Power LLC",
                  "Example 138kV", "Taylor", "WEST", datetime(2027, 6, 30), "GAS", "GT", mw, None, None,
                  datetime(2024, 9, 2), datetime(2024, 11, 1), None, None, "No", datetime(2025, 3, 3)]
                 + [None] * 9 + [datetime(2026, 8, 14), None])
    large.append(["21INR0012", "Solar Field", "SS Completed, FIS Started, No IA", "Sun Co LLC", "Sub 345kV", "Pecos",
                  "WEST", datetime(2026, 12, 1), "SOL", "PV", 250, None, None, datetime(2021, 2, 1)] + [None] * 16)
    large.append(["15INR0064b", "Split Wind B", "SS Completed, FIS Completed, IA", "Wind Co", "Sub 2", "Glasscock",
                  "WEST", datetime(2026, 12, 3), "WIN", "WT", 162.1] + [None] * 19)
    large.append(["Note: projects are listed until commercial operation."])
    small = wb.create_sheet(gis.SMALL)
    small.append(["INR", "Project Name", "Model Ready Date", "Interconnecting Entity", "POI Location", "County",
                  "CDR Reporting Zone", "Projected COD", "Fuel", "Technology", "Capacity (MW)",
                  "Change indicators: Proj Name, MW", "IA Signed", "Financial Security ", "Approved for Energization",
                  "Approved for Synchronization", "Comment"])
    small.append(["25INR0101", "Small Battery", None, "Store LLC", "Feeder 12", "Bexar", "SOUTH",
                  datetime(2026, 5, 1), "OTH", "BA", 9.9] + [None] * 6)
    inactive = wb.create_sheet(gis.INACTIVE)
    inactive.append(["INR", "Project Name", "County", "Fuel", "MW **", "Size Category", "Inactive Date"])
    inactive.append(["22INR0300", "Paused Wind", "Nolan", "WIN", 200, "Large", datetime(2025, 7, 1)])
    cancelled = wb.create_sheet(gis.CANCELLED)
    cancelled.append(["INR", "Project Name", "County", "Fuel", "MW **", "Size Category", "Cancel Date"])
    cancelled.append(["21INR0012", "Solar Field", "Pecos", "SOL", 250, "Large", datetime(2026, 9, 15)])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_parses_the_gis_sheets():
    rows = {e.queue_id: e.values for e in gis.parse_workbook(gis_workbook())}
    assert set(rows) == {"24INR0459", "21INR0012", "15INR0064b", "25INR0101", "22INR0300"}
    gas = rows["24INR0459"]
    assert (gas["name"], gas["developer"], gas["county"], gas["state"], gas["poi"], gas["zone"]) == (
        "Test Peaker", "Example Power LLC", "Taylor", "TX", "Example 138kV", "WEST")
    assert (gas["fuel"], gas["mw_requested"], gas["proposed_cod"], gas["queue_date"]) == (
        "Gas", 360.5, "2027-06-30", date(2024, 9, 2))
    assert gas["project_type"] == "Large Gen (Combustion (gas) Turbine, but not part of a Combined-Cycle)"
    assert (gas["status"], gas["status_class"]) == ("Approved for Energization", "under_construction")
    assert gas["study_phase"] == ("SS Completed, FIS Completed, IA; IA signed 2025-03-03; "
                                  "approved for energization 2026-08-14")
    assert rows["21INR0012"]["status"] == "Cancelled" and rows["21INR0012"]["withdrawn_date"] == date(2026, 9, 15)
    assert rows["25INR0101"]["project_type"] == "Small Gen (Battery Energy Storage)"
    assert rows["25INR0101"]["status"] == "In study"
    paused = rows["22INR0300"]
    assert (paused["status"], paused["status_class"], paused["mw_requested"]) == ("Inactive", "suspended", 200)
    assert gis._inr("15inr0064B") == "15INR0064b" and gis._inr("NOTES:") is None  # ERCOT's (and LBNL's) form


def test_report_month_spellings():
    assert gis.report_month("GIS_Report_September2026") == date(2026, 9, 1)
    assert gis.report_month("GIS_Report_Jun2026.xlsx") == date(2026, 6, 1)  # as posted 2026-07-01
    assert gis.report_month("GIS_Report_May_2020.xlsx") == date(2020, 5, 1)
    assert gis.report_month("Co-located_Battery_Identification_Report_September_2026") is None
    assert gis.report_month("GIS_Report_Ju2026") is None


def test_refuses_unknown_layouts():
    with pytest.raises(ParseError, match="not an xlsx"):
        gis.parse_workbook(b"<html>Access denied</html>")
    wb = openpyxl.Workbook()
    wb.active.title = "Summary"
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(ParseError, match="lacks sheets"):
        gis.parse_workbook(buf.getvalue())


def test_mis_listing_parses_the_captured_page():
    postings = gis.mis_listing(LISTING)
    assert len(postings) == 173 and sum(p.is_gis for p in postings) == 100
    first = next(p for p in postings if p.is_gis)
    assert first == gis.Posting(SEP_2026, "GIS_Report_September2026.xlsx", datetime(2026, 10, 1, 16, 8, 29), "15933")
    assert (first.name, first.posted, first.month) == ("GIS_Report_September2026", date(2026, 10, 1), date(2026, 9, 1))
    assert postings[0].file.startswith("Co-located_Battery") and not postings[0].is_gis
    assert all(p.month is not None for p in postings if p.is_gis)
    assert gis.mis_listing(b"<html><body>Maintenance</body></html>") == []


def test_plan_backfills_twelve_months_then_only_newer_reports():
    postings = gis.mis_listing(LISTING)
    todo, skipped = ercot.plan(postings, {})
    assert [p.name for p in todo] == [f"GIS_Report_{m}" for m in (
        "October2025", "November2025", "December2025", "January2026", "February2026", "March2026", "April2026",
        "May2026", "Jun2026", "July2026", "August2026", "September2026")] and skipped == []
    sep = todo[-1]
    assert ercot.plan(postings, {p.doc_id: p.posted_at for p in todo}) == ([], [])
    todo2, skipped2 = ercot.plan(postings, {sep.doc_id: sep.posted_at})  # never loaded out of order
    assert todo2 == [] and len(skipped2) == 11


def test_robots_terms_and_owner_decision():
    robots = (DATA / "robots.txt").read_text()
    assert "User-agent: *" in robots and "Disallow: /misapp" in robots and "Disallow: /misdownload" in robots
    assert all(line in ercot.TERMS_NOTE for line in ("User-agent: *", "Disallow: /misapp", "Disallow: /misdownload"))
    for rule in ("Owner's decision", 'User-Agent "Marden research cody@marden.ai"', "at most 1 request a second",
                 "weekly check", "only reports not already in the document table", "never evaded"):
        assert rule in ercot.TERMS_NOTE
    assert ercot.USER_AGENT == "Marden research cody@marden.ai" and ercot.PER_SECOND <= 1
    product = (DATA / "data_product_PG7-200-ER_excerpt.txt").read_text()
    assert f"EMIL ID pg7-200-er Report Type ID {ercot.REPORT_TYPE_ID}" in product and "Chron - Monthly" in product
    terms = " ".join((DATA / "terms_of_use_excerpt.txt").read_text().split())
    reuse = ("raw data provided in public portions of this website may be used, reproduced, and redistributed in "
             "compilations, charts, and analyses without maintaining such notices.")
    assert reuse in terms and reuse in ercot.TERMS_NOTE
    assert licensing.licence_for(ercot.SOURCE_NAME)["status"] == "public" and not ercot.SOURCE.blocked
    assert ercot.SOURCE.base_url == ercot.LIST_URL and ercot.SOURCE.kind == "portal"


@pytest.fixture
def fast_rate() -> Iterator[None]:
    http.set_rate_limit(ercot.HOST, 1000.0)
    yield
    http.set_rate_limit(ercot.HOST, ercot.PER_SECOND)


def _without_september(listing: bytes) -> bytes:
    """The captured listing as it stood before ERCOT posted the September 2026 report on 2026-10-01."""
    row = re.search(rb"<tr><td class='labelOptional_ind'>RPT[^<]*GIS_Report_September2026.*?</tr>", listing, re.S)
    assert row is not None
    return listing.replace(row[0], b"", 1)


class Ercot:
    """The MIS as captured, with a distinct workbook per posting; `stop` makes one URL refuse."""

    def __init__(self, listing: bytes = LISTING):
        self.listing, self.calls, self.stop = listing, [], None  # type: ignore[var-annotated]
        self.agents: set[str] = set()
        self.ids = {p.doc_id: i for i, p in enumerate(reversed(gis.mis_listing(LISTING)))}

    def handler(self, request: httpx.Request) -> httpx.Response:
        url = str(request.url)
        self.calls.append(url)
        self.agents.add(request.headers["user-agent"])
        if self.stop and self.stop[0] in url:
            return self.stop[1]()
        if url == ercot.LIST_URL:
            return httpx.Response(200, content=self.listing, headers={"content-type": "text/html;charset=ISO-8859-1"})
        doc_id = request.url.params.get("doclookupId")
        if request.url.path == "/misdownload/servlets/mirDownload" and doc_id in self.ids:
            return httpx.Response(200, content=gis_workbook(mw=300 + self.ids[doc_id]))
        return httpx.Response(404)

    def fetcher(self) -> ercot.Fetcher:
        return ercot.make_fetcher(httpx.Client(transport=httpx.MockTransport(self.handler)))

    def downloads(self) -> int:
        return sum("mirDownload" in c for c in self.calls)


def test_run_backfills_twelve_reports_once_and_reruns_clean(committed, fast_rate, monkeypatch):
    monkeypatch.setattr(ercot, "MIN_ENTRIES", 1)
    site = Ercot(_without_september(LISTING))
    first = ercot.run(fetcher=site.fetcher())
    assert first["ok"], first
    info = first["ercot_gis"]
    assert info["listed"] == 172 and info["gis_listed"] == 99 and info["already_loaded"] == 0
    assert [r["name"] for r in info["loaded"]][::11] == ["GIS_Report_September2025", "GIS_Report_August2026"]
    assert site.downloads() == 12 == len(info["downloaded"]) and info["requests"] == 13
    assert site.agents == {"Marden research cody@marden.ai"}
    assert site.calls[0] == ercot.LIST_URL and all("mirDownload?doclookupId=" in u for u in site.calls[1:])
    assert first["documents_new"] == 13 and first["iso"]["ercot"]["entries"] == 5
    assert committed.execute("SELECT count(*) AS n FROM queue_project WHERE iso = 'ercot'").fetchone()["n"] == 5

    second = ercot.run(fetcher=site.fetcher())
    assert second["ok"], second
    assert second["facts_new"] == 0 and second["documents_new"] == 0 and second["ercot_gis"]["loaded"] == []
    assert site.downloads() == 12 and second["ercot_gis"]["requests"] == 1  # the listing only

    site.listing = LISTING  # ERCOT posts September 2026
    third = ercot.run(fetcher=site.fetcher())
    assert third["ok"], third
    assert [r["name"] for r in third["ercot_gis"]["loaded"]] == ["GIS_Report_September2026"]
    assert site.downloads() == 13 and third["documents_new"] == 2  # the changed listing and the new report
    doc = documents.latest_document(ercot.SOURCE_ID, ercot.DOWNLOAD_URL.format(SEP_2026))
    assert doc is not None and doc.external_id == SEP_2026 and doc.doc_date == date(2026, 10, 1)
    assert doc.doc_type == "gis_xlsx" and doc.meta["report_month"] == "2026-09-01"
    mw = committed.execute("SELECT valid_from, source_url, tag, source_name FROM fact_current "
                           "WHERE subject_id = 'ercot:24INR0459' AND field = 'mw_requested'").fetchone()
    assert (mw["valid_from"], mw["source_url"], mw["tag"], mw["source_name"]) == (
        date(2026, 10, 1), doc.url, "F", "ERCOT")
    history = committed.execute("SELECT count(*) AS n FROM fact WHERE subject_id = 'ercot:24INR0459' "
                                "AND field = 'mw_requested'").fetchone()["n"]
    assert history == 13  # one value per monthly report (the MW differs in each), oldest first
    fourth = ercot.run(fetcher=site.fetcher())
    assert fourth["ok"] and fourth["facts_new"] == 0 and site.downloads() == 13


def _refuse(status: int, body: bytes = b"") -> httpx.Response:
    return httpx.Response(status, content=body, headers={"content-type": "text/html"})


@pytest.mark.parametrize(("where", "answer", "why"), [
    ("GetReports.do", lambda: _refuse(403, INCAPSULA.read_bytes()), "HTTP 403"),
    ("GetReports.do", lambda: _refuse(200, INCAPSULA.read_bytes()), "challenge page"),
    ("doclookupId=1210876874", lambda: _refuse(200, b"<html>Please solve this CAPTCHA</html>"), "challenge page"),
    ("doclookupId=1210876874", lambda: _refuse(429), "HTTP 429"),
])
def test_a_refusal_stops_the_run(committed, fast_rate, monkeypatch, where, answer, why):
    monkeypatch.setattr(ercot, "MIN_ENTRIES", 1)
    site = Ercot()
    site.stop = (where, answer)
    out = ercot.run(fetcher=site.fetcher())
    assert not out["ok"] and "Stopped" in out["error"] and why in out["error"], out
    assert sum(where in c for c in site.calls) == 1  # asked once, never retried
    assert where in site.calls[-1]  # nothing after it
    src = sources.get_source(ercot.SOURCE_ID)
    assert src is not None and why in src["last_error"]
    if where.startswith("doclookupId"):  # March 2026: the reports before it stay loaded
        assert [r["name"] for r in out["ercot_gis"]["loaded"]][-1] == "GIS_Report_February2026"
        again = Ercot()
        resumed = ercot.run(fetcher=again.fetcher())
        assert resumed["ok"] and resumed["ercot_gis"]["loaded"][0]["name"] == "GIS_Report_March2026"
    else:
        assert committed.execute("SELECT count(*) AS n FROM fact").fetchone()["n"] == 0


class _Clock:
    """A fake clock for marden.core.http alone: sleeping advances it, nothing really waits."""

    def __init__(self) -> None:
        self.now = 1000.0

    def monotonic(self) -> float:
        return self.now

    def sleep(self, seconds: float) -> None:
        self.now += seconds


def test_requests_to_www_ercot_com_are_five_seconds_apart(committed, monkeypatch):
    clock = _Clock()
    monkeypatch.setattr(http, "time", clock)
    monkeypatch.setattr(ercot, "MIN_ENTRIES", 1)
    http.set_rate_limit(ercot.HOST, ercot.PER_SECOND)  # a fresh bucket on the fake clock, at the real rate
    site, sent = Ercot(_without_september(LISTING)), []

    def timed(request: httpx.Request) -> httpx.Response:
        sent.append(clock.now)
        return site.handler(request)

    try:
        out = ercot.run(fetcher=ercot.make_fetcher(httpx.Client(transport=httpx.MockTransport(timed))))
        rate = http.bucket_for(ercot.HOST)[1].rate  # the host-wide bucket the `ercot` collector shares
    finally:
        http.set_rate_limit(ercot.HOST, ercot.PER_SECOND)  # drop the fake-clock bucket
    assert out["ok"], out
    assert len(sent) == 13 and rate == ercot.PER_SECOND == 0.2
    assert min(b - a for a, b in itertools.pairwise(sent)) >= 5.0 - 1e-9
    assert sent[-1] - sent[0] >= 12 * 5.0 - 1e-9


def _stored(name: str, doc_id: str, posted: datetime) -> tuple[documents.DocumentRow, gis.Posting]:
    doc = documents.record_document(ercot.SOURCE_ID, ercot.DOWNLOAD_URL.format(doc_id), (DATA / name).read_bytes(),
                                    content_type=ercot.XLSX, doc_type="gis_xlsx", doc_date=posted.date(),
                                    external_id=doc_id)
    return doc, gis.Posting(doc_id, name, posted, "15933")


def test_real_reports_load_in_order_and_link_goodnight(committed):
    """ERCOT's August and September 2026 reports (trimmed): history moves forward and the one Texas match
    Marden has (Goodnight Wind II, FGE Goodnight II LLC, Armstrong County, to campus `goodnight`) links."""
    committed.execute("INSERT INTO campus (id, name, county, state) "
                      "VALUES ('goodnight', 'Goodnight', 'Armstrong', 'TX')")
    sources.upsert_spec(ercot.SOURCE)
    aug = _stored("GIS_Report_August2026_trimmed.xlsx", "1269363208", datetime(2026, 9, 1, 14, 38, 5))
    sep = _stored("GIS_Report_September2026_trimmed.xlsx", "1281327706", datetime(2026, 10, 1, 16, 8, 29))
    first = ercot.load(*aug, ercot.RunStats(), min_entries=1)
    assert first["entries"] == 31 and first["linked"] == {"ercot:23INR0637": ["campus:goodnight"]}
    stats = ercot.RunStats()
    second = ercot.load(*sep, stats, min_entries=1)
    assert second["entries"] == 30 and second["entries_new"] == 5 and second["facts_new"] > 0
    cur = {r["field"]: r for r in committed.execute(
        "SELECT field, value, valid_from FROM fact_current WHERE subject_id = 'ercot:23INR0637'").fetchall()}
    assert {f: cur[f]["value"] for f in ("name", "developer", "county", "state", "mw_requested", "proposed_cod",
                                         "status", "zone")} == {
        "name": "Goodnight Wind II", "developer": "FGE Goodnight II, LLC", "county": "Armstrong", "state": "TX",
        "mw_requested": 261.1, "proposed_cod": "2027-06-06", "status": "IA Signed", "zone": "PANHANDLE"}
    assert cur["linked_campus"]["value"][0]["evidence"]["phrase"] == "goodnight"
    energized = {r["field"]: (r["value"], r["valid_from"]) for r in committed.execute(
        "SELECT field, value, valid_from FROM fact_current WHERE subject_id = 'ercot:21INR0253'").fetchall()}
    assert energized["status"] == ("Approved for Energization", date(2026, 10, 1))  # IA Signed in August
    assert energized["proposed_cod"] == ("2027-03-17", date(2026, 10, 1))
    n = committed.execute("SELECT count(*) AS n FROM fact WHERE subject_id = 'ercot:21INR0253' "
                          "AND field = 'status'").fetchone()["n"]
    assert n == 2
    cancelled = committed.execute("SELECT value FROM fact_current WHERE subject_id = 'ercot:20INR0029' "
                                  "AND field = 'withdrawn_date'").fetchone()["value"]
    assert cancelled == "2026-09-14"
    again = ercot.RunStats()
    ercot.load(*sep, again, min_entries=1)
    assert again.facts_new == 0 and again.facts_seen == stats.facts_seen
    with pytest.raises(ParseError, match="truncated"):
        ercot.load(*sep, ercot.RunStats())  # 30 entries is not a whole GIS report


def test_schedule_and_cli():
    assert app.tasks["collect.ercot_gis"].queue == "collect"
    crons = {name: t.cron for (name, _), t in app.periodic_registry.periodic_tasks.items()}
    assert crons["collect.ercot_gis"] == ercot.CRON_UTC == "17 13 * * 2"
    assert cli.build_parser().parse_args(["collect", "ercot_gis"])._handler is ercot.collect_cli


@pytest.mark.network
def test_live_listing_still_lists_gis_reports():
    """One request: the MIS listing still answers with GIS_Report postings and no challenge."""
    with ercot.make_fetcher() as fetcher:
        resp = fetcher.get(ercot.LIST_URL, headers={"User-Agent": ercot.USER_AGENT}, use_cache=False)
    assert not ercot.challenged(resp.content)
    gis_reports = [p for p in gis.mis_listing(resp.content) if p.is_gis]
    assert len(gis_reports) >= ercot.BACKFILL and all(p.month for p in gis_reports[:ercot.BACKFILL])
